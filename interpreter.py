"""
IndLan Interpreter
A tree-walking interpreter that executes the AST produced by the parser,
with a complete Hindi + English Python ecosystem bridge.
"""

import sys
import os
import csv
import json
from ast_nodes import *
from py_bridge import (
    import_module_dynamic,
    resolve_from_import,
    resolve_attribute,
    PyBridgeCallable,
    IndLanImportError,
    HINDI_METHOD_ALIASES,
)


class IndLanRuntimeError(Exception):
    def __init__(self, message, line=None):
        loc = f"[Line {line}] " if line is not None else ""
        super().__init__(f"{loc}Runtime Error: {message}")
        self.line = line


class BreakSignal(Exception):
    pass

class ContinueSignal(Exception):
    pass

class ReturnSignal(Exception):
    def __init__(self, value):
        self.value = value


class Environment:
    """A single lexical scope, chained to its parent (closures!)."""

    def __init__(self, parent=None):
        self.vars = {}
        self.parent = parent

    def define(self, name, value):
        self.vars[name] = value

    def get(self, name, line=None):
        env = self
        while env is not None:
            if name in env.vars:
                return env.vars[name]
            env = env.parent
        raise IndLanRuntimeError(f"Undefined variable '{name}'", line)

    def set(self, name, value, line=None):
        env = self
        while env is not None:
            if name in env.vars:
                env.vars[name] = value
                return
            env = env.parent
        self.vars[name] = value

    def has(self, name):
        env = self
        while env is not None:
            if name in env.vars:
                return True
            env = env.parent
        return False


class IndFunction:
    """A user-defined function (also used for class methods)."""

    def __init__(self, name, params, body, closure, interpreter, this_binding=None):
        self.name = name
        self.params = params
        self.body = body
        self.closure = closure
        self.interpreter = interpreter
        self.this_binding = this_binding  # for bound methods

    def bind(self, instance):
        return IndFunction(self.name, self.params, self.body, self.closure, self.interpreter, this_binding=instance)

    def call(self, pos_args, kw_args=None, line=None):
        kw_args = kw_args or {}
        env = Environment(self.closure)
        if self.this_binding is not None:
            env.define("this", self.this_binding)

        bound_params = set()
        # Bind positional args
        if len(pos_args) > len(self.params):
            raise IndLanRuntimeError(
                f"Function '{self.name}' expects at most {len(self.params)} positional argument(s), got {len(pos_args)}", line
            )
        for param, arg in zip(self.params, pos_args):
            env.define(param, arg)
            bound_params.add(param)

        # Bind keyword args
        for k, v in kw_args.items():
            if k in bound_params:
                raise IndLanRuntimeError(f"Function '{self.name}' got multiple values for argument '{k}'", line)
            if k not in self.params:
                raise IndLanRuntimeError(f"Function '{self.name}' got unexpected keyword argument '{k}'", line)
            env.define(k, v)
            bound_params.add(k)

        if len(bound_params) != len(self.params):
            missing = [p for p in self.params if p not in bound_params]
            raise IndLanRuntimeError(
                f"Function '{self.name}' missing required argument(s): {', '.join(missing)}", line
            )

        try:
            self.interpreter.exec_block(self.body, env)
        except ReturnSignal as r:
            return r.value
        return None

    def __repr__(self):
        return f"<function {self.name}>"


class IndClass:
    def __init__(self, name, methods):
        self.name = name
        self.methods = methods  # dict: name -> IndFunction (unbound)

    def find_method(self, name):
        return self.methods.get(name)

    def __repr__(self):
        return f"<class {self.name}>"


class IndInstance:
    def __init__(self, klass):
        self.klass = klass
        self.fields = {}

    def get(self, name, line=None):
        if name in self.fields:
            return self.fields[name]
        method = self.klass.find_method(name)
        if method is not None:
            return method.bind(self)
        raise IndLanRuntimeError(f"'{self.klass.name}' has no property '{name}'", line)

    def set(self, name, value):
        self.fields[name] = value

    def __repr__(self):
        return f"<{self.klass.name} instance>"


class Interpreter:
    def __init__(self, debug=False):
        self.debug = debug
        self.globals = Environment()
        self.setup_builtins()

    # ---------- built-ins ----------

    def setup_builtins(self):
        def ind_print(*args, **kwargs):
            print(*[self.stringify(a) for a in args])
            return None

        def ind_len(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if len(args) != 1:
                raise IndLanRuntimeError("len() expects exactly 1 argument", line)
            val = args[0]
            try:
                return len(val)
            except TypeError:
                raise IndLanRuntimeError(f"len() not supported for {type(val).__name__}", line)

        def ind_range(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if len(args) == 1:
                return list(range(int(args[0])))
            if len(args) == 2:
                return list(range(int(args[0]), int(args[1])))
            if len(args) == 3:
                return list(range(int(args[0]), int(args[1]), int(args[2])))
            raise IndLanRuntimeError("range() expects 1 to 3 arguments", line)

        def ind_str(*args, **kwargs):
            return self.stringify(args[0]) if args else ""

        def ind_int(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            try:
                return int(args[0])
            except (ValueError, TypeError, IndexError):
                raise IndLanRuntimeError(f"Cannot convert to int: {args[0] if args else ''!r}", line)

        def ind_float(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            try:
                return float(args[0])
            except (ValueError, TypeError, IndexError):
                raise IndLanRuntimeError(f"Cannot convert to float: {args[0] if args else ''!r}", line)

        def ind_input(*args, **kwargs):
            prompt = self.stringify(args[0]) if args else ""
            return input(prompt)

        def ind_type(*args, **kwargs):
            if not args:
                return "null"
            v = args[0]
            if isinstance(v, bool):
                return "bool"
            if isinstance(v, int):
                return "int"
            if isinstance(v, float):
                return "float"
            if isinstance(v, str):
                return "string"
            if isinstance(v, list):
                return "list"
            if isinstance(v, dict):
                return "dict"
            if v is None:
                return "null"
            if isinstance(v, IndInstance):
                return v.klass.name
            return type(v).__name__

        def ind_append(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if len(args) < 2:
                raise IndLanRuntimeError("append() expects a list and a value", line)
            lst, val = args[0], args[1]
            if not isinstance(lst, list):
                if hasattr(lst, "append"):
                    return lst.append(val)
                raise IndLanRuntimeError("append() expects a list as first argument", line)
            lst.append(val)
            return lst

        def ind_pop(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if not args:
                raise IndLanRuntimeError("pop() expects a list", line)
            lst = args[0]
            if not isinstance(lst, list):
                if hasattr(lst, "pop"):
                    return lst.pop(*args[1:])
                raise IndLanRuntimeError("pop() expects a list", line)
            if len(args) > 1:
                return lst.pop(int(args[1]))
            return lst.pop()

        def ind_keys(*args, **kwargs):
            if not args:
                return []
            d = args[0]
            return list(d.keys()) if hasattr(d, "keys") else []

        def ind_values(*args, **kwargs):
            if not args:
                return []
            d = args[0]
            return list(d.values()) if hasattr(d, "values") else []

        # Native CSV Functions
        def ind_csv_padho(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if not args and "filepath" not in kwargs:
                raise IndLanRuntimeError("csv_padho() requires a file path argument", line)
            path = args[0] if args else kwargs["filepath"]
            encoding = kwargs.get("encoding", "utf-8")
            delimiter = kwargs.get("delimiter", kwargs.get("sep", ","))
            try:
                with open(path, "r", encoding=encoding, newline="") as f:
                    reader = csv.DictReader(f, delimiter=delimiter)
                    return [dict(row) for row in reader]
            except Exception as e:
                raise IndLanRuntimeError(f"csv_padho failed to read '{path}': {e}", line)

        def ind_csv_likho(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if len(args) < 2 and ("filepath" not in kwargs or "data" not in kwargs):
                raise IndLanRuntimeError("csv_likho() requires filepath and data arguments", line)
            path = args[0] if len(args) > 0 else kwargs["filepath"]
            data = args[1] if len(args) > 1 else kwargs["data"]
            encoding = kwargs.get("encoding", "utf-8")
            delimiter = kwargs.get("delimiter", kwargs.get("sep", ","))

            if hasattr(data, "to_csv"):
                try:
                    data.to_csv(path, index=kwargs.get("index", False), sep=delimiter, encoding=encoding)
                    return None
                except Exception as e:
                    raise IndLanRuntimeError(f"csv_likho failed to write DataFrame to '{path}': {e}", line)
            try:
                with open(path, "w", encoding=encoding, newline="") as f:
                    if not data:
                        return None
                    if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
                        fieldnames = list(data[0].keys())
                        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=delimiter)
                        writer.writeheader()
                        writer.writerows(data)
                    elif isinstance(data, list) and len(data) > 0 and isinstance(data[0], (list, tuple)):
                        writer = csv.writer(f, delimiter=delimiter)
                        writer.writerows(data)
                    else:
                        for item in data:
                            f.write(str(item) + "\n")
                return None
            except Exception as e:
                raise IndLanRuntimeError(f"csv_likho failed to write '{path}': {e}", line)

        def ind_csv_jodo(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if len(args) < 2:
                raise IndLanRuntimeError("csv_jodo() expects filepath and row", line)
            path, row = args[0], args[1]
            encoding = kwargs.get("encoding", "utf-8")
            delimiter = kwargs.get("delimiter", kwargs.get("sep", ","))
            try:
                with open(path, "a", encoding=encoding, newline="") as f:
                    if isinstance(row, dict):
                        writer = csv.DictWriter(f, fieldnames=list(row.keys()), delimiter=delimiter)
                        writer.writerow(row)
                    elif isinstance(row, (list, tuple)):
                        writer = csv.writer(f, delimiter=delimiter)
                        writer.writerow(row)
                    else:
                        f.write(str(row) + "\n")
                return None
            except Exception as e:
                raise IndLanRuntimeError(f"csv_jodo failed: {e}", line)

        def ind_csv_badlo(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if len(args) < 2:
                raise IndLanRuntimeError("csv_badlo() expects filepath and transform function", line)
            path, fn = args[0], args[1]
            rows = ind_csv_padho(path, **kwargs)
            transformed = []
            for r in rows:
                if isinstance(fn, IndFunction):
                    new_r = fn.call([r], {}, line)
                elif callable(fn):
                    new_r = fn(r)
                else:
                    raise IndLanRuntimeError("csv_badlo transform must be callable", line)
                transformed.append(new_r if new_r is not None else r)
            ind_csv_likho(path, transformed, **kwargs)
            return transformed

        def ind_csv_chhano(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if len(args) < 2:
                raise IndLanRuntimeError("csv_chhano() expects filepath and filter predicate", line)
            path, fn = args[0], args[1]
            rows = ind_csv_padho(path, **kwargs)
            filtered = []
            for r in rows:
                if isinstance(fn, IndFunction):
                    keep = fn.call([r], {}, line)
                elif callable(fn):
                    keep = fn(r)
                else:
                    raise IndLanRuntimeError("csv_chhano predicate must be callable", line)
                if keep:
                    filtered.append(r)
            return filtered

        # Native JSON Functions
        def ind_json_padho(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if not args and "filepath" not in kwargs:
                raise IndLanRuntimeError("json_padho() requires a file path argument", line)
            path = args[0] if args else kwargs["filepath"]
            try:
                with open(path, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception as e:
                raise IndLanRuntimeError(f"json_padho failed to read '{path}': {e}", line)

        def ind_json_likho(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if len(args) < 2 and ("filepath" not in kwargs or "data" not in kwargs):
                raise IndLanRuntimeError("json_likho() requires filepath and data arguments", line)
            path = args[0] if len(args) > 0 else kwargs["filepath"]
            data = args[1] if len(args) > 1 else kwargs["data"]
            indent = kwargs.get("indent", 2)
            try:
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(data, f, indent=indent, default=str)
                return None
            except Exception as e:
                raise IndLanRuntimeError(f"json_likho failed to write '{path}': {e}", line)

        # Native Excel Functions
        def ind_excel_padho(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if not args and "filepath" not in kwargs:
                raise IndLanRuntimeError("excel_padho() requires a file path argument", line)
            path = args[0] if args else kwargs["filepath"]
            try:
                import pandas as pd
                return pd.read_excel(path, **kwargs)
            except ImportError:
                pass
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, data_only=True)
                sheet = wb.active
                data = []
                headers = [cell.value for cell in next(sheet.iter_rows())]
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    data.append(dict(zip(headers, row)))
                return data
            except ImportError:
                raise IndLanRuntimeError(
                    "[IndLan Import Error]\nNeither 'pandas' nor 'openpyxl' is installed.\nInstall using: pip install openpyxl",
                    line
                )
            except Exception as e:
                raise IndLanRuntimeError(f"excel_padho failed: {e}", line)

        def ind_excel_likho(*args, **kwargs):
            line = kwargs.pop("__line__", None)
            if len(args) < 2 and ("filepath" not in kwargs or "data" not in kwargs):
                raise IndLanRuntimeError("excel_likho() requires filepath and data arguments", line)
            path = args[0] if len(args) > 0 else kwargs["filepath"]
            data = args[1] if len(args) > 1 else kwargs["data"]
            if hasattr(data, "to_excel"):
                try:
                    data.to_excel(path, index=kwargs.get("index", False))
                    return None
                except Exception as e:
                    raise IndLanRuntimeError(f"excel_likho failed: {e}", line)
            try:
                import pandas as pd
                if isinstance(data, list):
                    df = pd.DataFrame(data)
                    df.to_excel(path, index=False)
                    return None
            except ImportError:
                pass
            try:
                import openpyxl
                wb = openpyxl.Workbook()
                ws = wb.active
                if data and isinstance(data[0], dict):
                    headers = list(data[0].keys())
                    ws.append(headers)
                    for item in data:
                        ws.append([item.get(h) for h in headers])
                elif data and isinstance(data[0], (list, tuple)):
                    for row in data:
                        ws.append(list(row))
                wb.save(path)
                return None
            except ImportError:
                raise IndLanRuntimeError(
                    "[IndLan Import Error]\nNeither 'pandas' nor 'openpyxl' is installed.\nInstall using: pip install openpyxl",
                    line
                )
            except Exception as e:
                raise IndLanRuntimeError(f"excel_likho failed: {e}", line)

        builtins = {
            "print": ind_print,
            "chhap": ind_print,
            "len": ind_len,
            "range": ind_range,
            "str": ind_str,
            "int": ind_int,
            "float": ind_float,
            "input": ind_input,
            "type": ind_type,
            "append": ind_append,
            "pop": ind_pop,
            "keys": ind_keys,
            "values": ind_values,
            "csv_padho": ind_csv_padho,
            "csv_likho": ind_csv_likho,
            "csv_jodo": ind_csv_jodo,
            "csv_badlo": ind_csv_badlo,
            "csv_chhano": ind_csv_chhano,
            "json_padho": ind_json_padho,
            "json_likho": ind_json_likho,
            "excel_padho": ind_excel_padho,
            "excel_likho": ind_excel_likho,
        }
        for name, fn in builtins.items():
            self.globals.define(name, fn)

    # ---------- driver ----------

    def run(self, program):
        env = self.globals
        for stmt in program.statements:
            self.execute(stmt, env)

    # ---------- statement execution ----------

    def execute(self, node, env):
        method = getattr(self, f"exec_{type(node).__name__}", None)
        if method is None:
            raise IndLanRuntimeError(f"No executor for statement {type(node).__name__}")
        return method(node, env)

    def exec_block(self, block, env):
        for stmt in block.statements:
            self.execute(stmt, env)

    def exec_Block(self, node, env):
        inner = Environment(env)
        self.exec_block(node, inner)

    def exec_ImportStmt(self, node, env):
        mod = import_module_dynamic(node.module_path, node.line)
        if node.alias:
            env.define(node.alias, mod)
        else:
            # e.g. import pandas -> define 'pandas'
            # e.g. import matplotlib.pyplot -> define 'matplotlib' and 'pyplot'
            parts = node.module_path.split(".")
            root_mod = import_module_dynamic(parts[0], node.line)
            env.define(parts[0], root_mod)
            if len(parts) > 1:
                env.define(parts[-1], mod)

    def exec_FromImportStmt(self, node, env):
        for item_name, alias in node.items:
            obj = resolve_from_import(node.module_path, item_name, node.line)
            target_name = alias if alias else item_name
            env.define(target_name, obj)

    def exec_LetStmt(self, node, env):
        value = self.evaluate(node.value, env)
        if len(node.names) == 1:
            env.define(node.names[0], value)
        else:
            try:
                items = list(value)
            except (TypeError, ValueError):
                raise IndLanRuntimeError(f"Cannot unpack non-iterable {type(value).__name__}", node.line)
            if len(items) != len(node.names):
                raise IndLanRuntimeError(
                    f"Cannot unpack {len(items)} values into {len(node.names)} variables", node.line
                )
            for name, item in zip(node.names, items):
                env.define(name, item)

    def exec_ExprStmt(self, node, env):
        self.evaluate(node.expr, env)

    def exec_FunDecl(self, node, env):
        fn = IndFunction(node.name, node.params, node.body, env, self)
        env.define(node.name, fn)

    def exec_ClassDecl(self, node, env):
        methods = {}
        for m in node.methods:
            methods[m.name] = IndFunction(m.name, m.params, m.body, env, self)
        klass = IndClass(node.name, methods)
        env.define(node.name, klass)

    def exec_IfStmt(self, node, env):
        for cond, body in node.branches:
            if self.truthy(self.evaluate(cond, env)):
                inner = Environment(env)
                self.exec_block(body, inner)
                return
        if node.else_block is not None:
            inner = Environment(env)
            self.exec_block(node.else_block, inner)

    def exec_WhileStmt(self, node, env):
        while self.truthy(self.evaluate(node.condition, env)):
            inner = Environment(env)
            try:
                self.exec_block(node.body, inner)
            except BreakSignal:
                break
            except ContinueSignal:
                continue

    def exec_DoWhileStmt(self, node, env):
        while True:
            inner = Environment(env)
            try:
                self.exec_block(node.body, inner)
            except BreakSignal:
                break
            except ContinueSignal:
                pass
            if not self.truthy(self.evaluate(node.condition, env)):
                break

    def exec_SwitchStmt(self, node, env):
        subject_val = self.evaluate(node.subject, env)
        for value_expr, body in node.cases:
            case_val = self.evaluate(value_expr, env)
            if self.is_equal(subject_val, case_val):
                inner = Environment(env)
                try:
                    self.exec_block(body, inner)
                except BreakSignal:
                    pass
                return
        if node.default_block is not None:
            inner = Environment(env)
            try:
                self.exec_block(node.default_block, inner)
            except BreakSignal:
                pass

    def exec_ForStmt(self, node, env):
        iterable = self.evaluate(node.iterable, env)
        if isinstance(iterable, dict):
            items = list(iterable.keys())
        elif isinstance(iterable, (list, str, tuple)):
            items = list(iterable)
        elif hasattr(iterable, "__iter__"):
            items = list(iterable)
        else:
            raise IndLanRuntimeError(f"Cannot iterate over this value", node.line)
        for item in items:
            inner = Environment(env)
            inner.define(node.var_name, item)
            try:
                self.exec_block(node.body, inner)
            except BreakSignal:
                break
            except ContinueSignal:
                continue

    def exec_ReturnStmt(self, node, env):
        value = self.evaluate(node.value, env) if node.value is not None else None
        raise ReturnSignal(value)

    def exec_BreakStmt(self, node, env):
        raise BreakSignal()

    def exec_ContinueStmt(self, node, env):
        raise ContinueSignal()

    # ---------- expression evaluation ----------

    def evaluate(self, node, env):
        method = getattr(self, f"eval_{type(node).__name__}", None)
        if method is None:
            raise IndLanRuntimeError(f"No evaluator for expression {type(node).__name__}")
        return method(node, env)

    def eval_NumberLit(self, node, env):
        return node.value

    def eval_StringLit(self, node, env):
        return node.value

    def eval_BoolLit(self, node, env):
        return node.value

    def eval_NullLit(self, node, env):
        return None

    def eval_ListLit(self, node, env):
        return [self.evaluate(e, env) for e in node.elements]

    def eval_DictLit(self, node, env):
        result = {}
        for k, v in node.pairs:
            key = self.evaluate(k, env)
            result[key] = self.evaluate(v, env)
        return result

    def eval_Identifier(self, node, env):
        return env.get(node.name, node.line)

    def eval_UnaryOp(self, node, env):
        val = self.evaluate(node.operand, env)
        if node.op == "-":
            self.check_number(val, node.line)
            return -val
        if node.op == "not":
            return not self.truthy(val)
        raise IndLanRuntimeError(f"Unknown unary operator '{node.op}'", node.line)

    def eval_BinOp(self, node, env):
        left = self.evaluate(node.left, env)
        right = self.evaluate(node.right, env)
        op = node.op
        line = node.line

        if op == "+":
            if isinstance(left, str) or isinstance(right, str):
                return self.stringify(left) + self.stringify(right)
            if isinstance(left, list) and isinstance(right, list):
                return left + right
            self.check_number(left, line); self.check_number(right, line)
            return left + right
        if op == "-":
            self.check_number(left, line); self.check_number(right, line)
            return left - right
        if op == "*":
            self.check_number(left, line); self.check_number(right, line)
            return left * right
        if op == "/":
            self.check_number(left, line); self.check_number(right, line)
            if right == 0:
                raise IndLanRuntimeError("Division by zero", line)
            result = left / right
            if isinstance(left, int) and isinstance(right, int) and left % right == 0:
                return int(result)
            return round(result, 10)
        if op == "%":
            self.check_number(left, line); self.check_number(right, line)
            if right == 0:
                raise IndLanRuntimeError("Modulo by zero", line)
            return left % right
        if op == "==":
            return self.is_equal(left, right)
        if op == "!=":
            return not self.is_equal(left, right)
        if op == "<":
            return left < right
        if op == ">":
            return left > right
        if op == "<=":
            return left <= right
        if op == ">=":
            return left >= right

        raise IndLanRuntimeError(f"Unknown binary operator '{op}'", line)

    def eval_LogicalOp(self, node, env):
        left = self.evaluate(node.left, env)
        if node.op == "and":
            if not self.truthy(left):
                return left
            return self.evaluate(node.right, env)
        else:  # or
            if self.truthy(left):
                return left
            return self.evaluate(node.right, env)

    def assign_single_target(self, target, value, op, line, env):
        if op != "=":
            current = self.evaluate(target, env)
            base_op = op[0]  # '+', '-', '*', '/'
            fake_bin = BinOp(base_op, Wrapped(current), Wrapped(value), line)
            value = self.eval_BinOp(fake_bin, env)

        if isinstance(target, Identifier):
            env.set(target.name, value, line)
        elif isinstance(target, Index):
            obj = self.evaluate(target.obj, env)
            idx = self.evaluate(target.index, env)
            if isinstance(obj, list):
                obj[int(idx)] = value
            elif isinstance(obj, dict):
                obj[idx] = value
            elif hasattr(obj, "__setitem__"):
                obj[idx] = value
            else:
                raise IndLanRuntimeError("Cannot index-assign to this type", line)
        elif isinstance(target, GetAttr):
            obj = self.evaluate(target.obj, env)
            if isinstance(obj, IndInstance):
                obj.set(target.name, value)
            else:
                setattr(obj, target.name, value)
        else:
            raise IndLanRuntimeError("Invalid assignment target", line)
        return value

    def eval_Assign(self, node, env):
        value = self.evaluate(node.value, env)
        if isinstance(node.target, list):
            try:
                items = list(value)
            except (TypeError, ValueError):
                raise IndLanRuntimeError(f"Cannot unpack non-iterable {type(value).__name__}", node.line)
            if len(items) != len(node.target):
                raise IndLanRuntimeError(
                    f"Cannot unpack {len(items)} values into {len(node.target)} targets", node.line
                )
            for t, item in zip(node.target, items):
                self.assign_single_target(t, item, node.op, node.line, env)
            return value
        return self.assign_single_target(node.target, value, node.op, node.line, env)

    def eval_Call(self, node, env):
        line = node.line
        pos_args = []
        kw_args = {}
        for a in node.args:
            if isinstance(a, KeywordArg):
                kw_args[a.name] = self.evaluate(a.value, env)
            else:
                pos_args.append(self.evaluate(a, env))

        # Check if callee is an IndClass instantiation
        if isinstance(node.callee, Identifier) and env.has(node.callee.name):
            callee_val = env.get(node.callee.name, line)
            if isinstance(callee_val, IndClass):
                instance = IndInstance(callee_val)
                init = callee_val.find_method("init")
                if init is not None:
                    init.bind(instance).call(pos_args, kw_args, line)
                return instance

        callee = self.evaluate(node.callee, env)

        if isinstance(callee, IndFunction):
            return callee.call(pos_args, kw_args, line)

        if isinstance(callee, PyBridgeCallable):
            try:
                return callee(*pos_args, **kw_args)
            except IndLanRuntimeError:
                raise
            except Exception as e:
                if self.debug:
                    import traceback
                    traceback.print_exc()
                raise IndLanRuntimeError(f"Error calling {callee.name}: {e}", line) from e

        if callable(callee):
            try:
                return callee(*pos_args, **kw_args)
            except IndLanRuntimeError:
                raise
            except Exception as e:
                if self.debug:
                    import traceback
                    traceback.print_exc()
                name = getattr(callee, "__name__", str(callee))
                raise IndLanRuntimeError(f"Error calling '{name}': {e}", line) from e

        raise IndLanRuntimeError(f"'{self.stringify(callee)}' is not callable", line)

    def eval_Index(self, node, env):
        obj = self.evaluate(node.obj, env)
        idx = self.evaluate(node.index, env)
        if isinstance(obj, list):
            i = int(idx)
            if i < -len(obj) or i >= len(obj):
                raise IndLanRuntimeError(f"List index {i} out of range", node.line)
            return obj[i]
        if isinstance(obj, dict):
            if idx not in obj:
                raise IndLanRuntimeError(f"Key {idx!r} not found in dict", node.line)
            return obj[idx]
        if isinstance(obj, str):
            i = int(idx)
            if i < -len(obj) or i >= len(obj):
                raise IndLanRuntimeError(f"String index {i} out of range", node.line)
            return obj[i]
        if hasattr(obj, "__getitem__"):
            try:
                return obj[idx]
            except Exception as e:
                raise IndLanRuntimeError(f"Index error on {type(obj).__name__}: {e}", node.line) from e

        raise IndLanRuntimeError(f"Cannot index object of type '{type(obj).__name__}'", node.line)

    def eval_GetAttr(self, node, env):
        obj = self.evaluate(node.obj, env)
        if isinstance(obj, IndInstance):
            return obj.get(node.name, node.line)

        # Python object/module/type attribute resolution
        found, val = resolve_attribute(obj, node.name, node.line)
        if found:
            if callable(val) and not isinstance(val, type):
                return PyBridgeCallable(val, node.name)
            return val

        raise IndLanRuntimeError(
            f"Method or property '{node.name}' could not be resolved on '{type(obj).__name__}'",
            node.line
        )

    # ---------- helpers ----------

    def check_number(self, val, line):
        if isinstance(val, bool) or not isinstance(val, (int, float)):
            raise IndLanRuntimeError(f"Expected a number, got {self.stringify(val)}", line)

    def truthy(self, val):
        if val is None or val is False:
            return False
        if val == 0:
            return False
        if val == "":
            return False
        if isinstance(val, (list, dict)) and len(val) == 0:
            return False
        if hasattr(val, "empty") and getattr(val, "empty", False) is True:
            return False
        return True

    def is_equal(self, a, b):
        if type(a) != type(b) and not (isinstance(a, (int, float)) and isinstance(b, (int, float))):
            return False
        try:
            eq = (a == b)
            if hasattr(eq, "all"):
                return bool(eq.all())
            return bool(eq)
        except Exception:
            return False

    def stringify(self, val):
        if val is None:
            return "null"
        if val is True:
            return "true"
        if val is False:
            return "false"
        if isinstance(val, float):
            if val == int(val):
                return f"{val:.1f}"
            return str(val)
        if isinstance(val, list):
            return "[" + ", ".join(self.stringify(v) for v in val) + "]"
        if isinstance(val, dict):
            items = ", ".join(f"{self.stringify(k)}: {self.stringify(v)}" for k, v in val.items())
            return "{" + items + "}"
        return str(val)


class Wrapped(Node):
    """Helper node used internally to feed already-evaluated values back into eval_BinOp."""
    def __init__(self, value):
        self.value = value
        self.line = 0


def _eval_wrapped(self, node, env):
    return node.value

Interpreter.eval_Wrapped = _eval_wrapped